#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""端到端示例：生成一份符合《本科生论文（设计）撰写要求》的论文骨架。

用法（把 <skill> 换成本 skill 目录）：
    <python> "<skill>/scripts/example_thesis.py" 论文初稿.docx

生成后请用 audit_thesis.py 校验，并在 Word 中按 F9 更新目录域。
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
import thesisdocx as T  # noqa: E402


def build(path, title="论文题目（小3号黑体）",
          subtitle=None, author_info=()):
    doc = T.new_thesis_document()

    # ------------------------------------------------------------------
    # 1) 封面 / 扉页 —— 第一节：不设页眉、不编页码
    #    （学院通常提供固定封面模板，这里只演示结构）
    # ------------------------------------------------------------------
    for _ in range(3):
        T.blank_line(doc)
    T.add_title(doc, title, subtitle)
    for _ in range(3):
        T.blank_line(doc)
    for line in author_info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        T.set_paragraph_spacing(p, T.LINE_SPACING)
        T._no_first_line_indent(p)
        T.set_run_font(p.add_run(line), T.PT_BODY)

    # ------------------------------------------------------------------
    # 2) 前置部分 —— 大写罗马数字页码，有页眉
    #    含：原创性及知识产权声明、开题报告表、目录、中文摘要、外文摘要
    # ------------------------------------------------------------------
    fm = T.start_front_matter_section(doc)
    T.set_header_text(fm, title)

    # 原创性及知识产权声明（各校文本不同，此处留说明）
    # 声明页题头不是正文章节标题，为避免被当作一级标题判定“上下未空一行”，
    # 这里用 add_front_matter_head + 居中正文的方式排，而不是 add_heading。
    T.add_front_matter_head(doc, "声明", T.PT_TOC_HEAD)
    T.add_body(doc, "本人郑重声明：所提交的论文（设计）是本人在导师指导下独立完成的成果。"
                    "（此处按学校提供的声明文本填写，本人签名、日期。）", indent=False)
    T.add_page_break(doc)

    # 中文摘要 —— 不写论文题目；“摘　要”小2加粗宋体，其后隔一行
    T.add_front_matter_head(doc, "摘要")
    T.add_body(doc, "本文针对……问题，提出了……方法，实验表明……。")
    T.add_keywords(doc, ["关键词一", "关键词二", "关键词三"])
    T.add_page_break(doc)

    # 外文摘要
    T.add_english_abstract_head(doc, "Abstract")
    T.add_english_body(doc, "This thesis addresses ... and proposes ... .")
    T.add_english_body(doc, "keyword one; keyword two", indent=False,
                       bold_lead="Keywords\u3000")
    T.add_page_break(doc)

    # 目录（域，需在 Word 中更新）
    T.add_toc_field(doc, "目　录")

    # ------------------------------------------------------------------
    # 3) 正文 —— 另起一节，阿拉伯数字页码从 1 开始
    # ------------------------------------------------------------------
    body = T.start_body_section(doc)
    T.set_header_text(body, "一、引言")   # 页眉=本页一级标题，换章节要改

    T.add_heading(doc, 1, "引言")
    T.add_body(doc, "随着……，该问题日益受到关注。本文的研究内容包括……。")
    T.add_heading(doc, 2, "研究背景与意义")
    T.add_body(doc, "国内外研究现状表明……。")
    T.add_heading(doc, 3, "国内研究现状")
    T.add_body(doc, "国内学者主要从……角度展开研究。")
    T.add_heading(doc, 4, "早期研究")
    T.add_body(doc, "早期研究以理论分析为主。")

    # 公式：居中、末尾不加标点；带编号时右对齐
    T.add_body(doc, "定义如下损失函数：")
    T.add_formula(doc, "L = -Σ yᵢ log ŷᵢ", number=1)

    # 三线表：表序表名在表上（先写题注再建表）
    T.add_table_caption(doc, 1, "实验数据统计")
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    rows = [["编号", "方法", "准确率/%", "召回率/%"],
            ["1", "基线", "85.2", "83.1"],
            ["2", "本文方法", "91.5", "90.2"],
            ["3", "——", "——", "——"]]        # 空缺格填 ——
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    T.apply_three_line_table(table)           # 上下1.5磅、表头下0.5磅、无左右边线
    T.style_table_text(table)                 # 表内文字小5号宋体
    T.add_table_note(doc, "注：数据来源于实测，单位符号置于表头右上角。")

    # 插图：图号图名在图下、居中、5号宋体；引用图加引用号上标
    T.add_body(doc, "系统结构如图1所示。")
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # pic.add_run().add_picture("system.png", width=Cm(10))   # 放入真实图片
    T.add_figure_caption(doc, 1, "系统总体结构", cite=1)
    T.add_table_note(doc, "注：坐标轴单位已在图中标明。")

    T.add_heading(doc, 1, "结论")
    T.add_body(doc, "本文完成了……，后续工作将……。")

    # ------------------------------------------------------------------
    # 4) 参考文献与致谢
    # ------------------------------------------------------------------
    T.add_heading(doc, 1, "参考文献")
    T.add_reference(doc, 1, "张三.数据结构[M].北京：高等教育出版社，2020.15-20.")
    T.add_reference(doc, 2, "李四.算法研究[J].计算机学报，2021，44（3）：1-9.")
    T.add_reference(doc, 3, "王五.在线资源[EB/OL].http://example.com，2023-01-01.")

    T.add_heading(doc, 1, "致谢")
    T.add_body(doc, "感谢导师的悉心指导。")

    doc.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "论文初稿.docx"
    build(out, author_info=("学　　院：计算机学院",
                            "专　　业：计算机科学与技术",
                            "学　　号：2021001",
                            "姓　　名：张三",
                            "指导教师：李四"))
    print("已生成:", os.path.abspath(out))
    print("下一步：")
    print("  1) 在 Word 中打开，右键目录 → 更新域 → 更新整个目录")
    print("  2) 运行 audit_thesis.py 校验")
